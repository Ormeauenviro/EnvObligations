"""Phase 3 transform: TP002.1 workbook -> clean per-list JSON and CSV for the SharePoint load.

Implements the mapping and normalisation approved in docs/data-audit.md
(decisions D-01 to D-16) against the schema in provision/list-schema.json.

Usage:
    python import/transform.py [--xlsx path] [--out import/output] [--import-date 2026-09-29]

Outputs (one pair per list, plus a review report):
    import/output/<ListKey>.json   loader input: key fields, live fields, items
    import/output/<ListKey>.csv    the same items flattened for review in Excel
    import/output/import-report.md what needs a human: unowned items, evidence gaps, review comments...
"""

from __future__ import annotations

import argparse
import csv
import datetime as dt
import json
import re
import sys
import warnings
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import openpyxl

sys.path.insert(0, str(Path(__file__).resolve().parent))
import tp002_layout as L  # noqa: E402

REPO = Path(__file__).resolve().parent.parent
DEFAULT_XLSX = REPO / "source" / "DTI-HSEQ-TP002.1_Environmental_Compliance_Obligations_Register_Rev2_DRAFT_DH.xlsx"
DEFAULT_OUT = REPO / "import" / "output"
CONFIG = REPO / "import" / "config"
SCHEMA = json.loads((REPO / "provision" / "list-schema.json").read_text(encoding="utf-8"))

EN_DASH = "–"
NA_INFO = f"N/A {EN_DASH} Info only"

# ---------------------------------------------------------------------------
# Obligation ID prefixes (D-15). Marine Plants blocks use the permit short name.
# ---------------------------------------------------------------------------
PREFIX = {
    "CEMP and General Requirements": "CEMP",
    "PFAS CMP Requirements": "PFAS",
    "ADR WWBW Conditions": "ADR",
    "Bridge Creek DA Conditions": "BC-DA",
    "No Name DA Conditions": "NN-DA",
    "External CHMP Conditions": "CHMP",
    "RPP Exemption": "RPP",
    "Specification Requirements": "SPEC",
    "Legislation": "LEG",
}

# ---------------------------------------------------------------------------
# Normalisation tables (docs/data-audit.md section 4)
# ---------------------------------------------------------------------------
STAGE_MAP = {
    "all": "All stages",
    "pre construction": "Pre-construction",
    "pre-construction": "Pre-construction",
    "pre and during construction": "Pre-construction and construction",
    "during construction": "Construction",
    "construction": "Construction",
    "construction / operation": "Construction and operation",
    "design / construction / operation": "Construction and operation",
    "construction / rehabilitation": "Construction and operation",
    "operation / maintenance": "Operation",
    "completion": "Completion",
    "decommissioning": "Completion",
    "pre-construction / completion": "Completion",
    "post construction": "Post-construction",
    "rehabilitation": "Rehabilitation and monitoring",
    "rehabilitation monitoring": "Rehabilitation and monitoring",
    "monitoring": "Rehabilitation and monitoring",
    "operational works": "Construction",  # Marine Plants: approval type, not a stage
    "temporary access": "Construction",
}

STATUS_MAP = {
    "": "Not started",
    "not completed": "Underway",
    "not completed - ongoing": "Underway",
    "not completed - rehabilitation": "Not started",
    "ongoing": "Underway",
    "completed": "Complete",
    "noted": NA_INFO,
}

# D-01b: Bridge Creek "Not completed" rows that are completion-triggered start as Not started.
STATUS_OVERRIDE = {
    ("Bridge Creek DA Conditions", 8): "Not started",   # Condition 3 - rehabilitation, "In Landscaping Scope"
    ("Bridge Creek DA Conditions", 13): "Not started",  # Condition 8 - restore profiles on completion
}

# Timing text -> (frequency, due anchor, anchor offset days). First matching pattern wins.
# Patterns run against lower-case text with whitespace collapsed and citations removed.
TIMING_RULES: list[tuple[str, str, str | None, int | None]] = [
    (r"^quarterly / six monthly\?", "Quarterly", None, None),           # CEMP row 41: reviewer's open question
    (r"^weekly monthly$", "Weekly", None, None),                        # CEMP row 11: split (D-05)
    (r"^daily weekly$", "Weekly", None, None),
    (r"^weekly check sampling during discharge", "Trigger-based", None, None),
    (r"^weekly", "Weekly", None, None),
    (r"^daily during", "Trigger-based", None, None),
    (r"^daily monitoring$", "Trigger-based", None, None),                # BC: only while dewatering
    (r"^daily$", "Daily", None, None),
    (r"^ad hoc monthly as part of reporting", "Monthly", None, None),
    (r"^monthly", "Monthly", None, None),
    (r"^ad hoc", "Trigger-based", None, None),
    (r"^quarterly after revegetation commences", "Quarterly", "Revegetation commencement", 0),
    (r"^quarterly", "Quarterly", None, None),
    (r"^annual", "Annual", None, None),
    (r"^prior to construction works commencing / at the commencement", "Trigger-based", None, None),
    (r"^prior to construction works commencing", "Once", "Construction start", 0),
    (r"^prior to clearing works", "Trigger-based", None, None),
    (r"^prior to disposal or release", "Trigger-based", None, None),
    (r"^post construction - within four weeks", "Once", "Practical completion", 28),
    (r"^post construction", "Once", "Practical completion", 0),
    (r"^(at all times|duration of works|for duration of works|construction|construction phase|"
     r"construction and operation|rehabilitation phase|ongoing|as specified in plans)$", "Ongoing", None, None),
    (r"^(as required|when fish stranding)", "Trigger-based", None, None),
    (r"^during (dewatering|pumping activities|salvage|planting)", "Trigger-based", None, None),
    (r"^(prior to dewatering|following salvage works)$", "Once", None, None),
    (r"^within 2 weeks of rehabilitation completion", "Once", "Rehabilitation completion", 14),
    (r"^within 10 business days of completion and before post-works", "Once", "Practical completion", 14),
    (r"^within 10 business days of removal", "Once", "Temporary barrier removal", 14),
    (r"^within 15 business days of completion", "Once", "Practical completion", 21),
    (r"^years 1, 2 and 5 post-construction", "Milestone", "Practical completion", None),
    (r"^(years 1, 2 and 5|monitoring milestones|reporting milestones)", "Milestone", None, None),
    (r"^5-20 business days before start", "Once", None, None),
    (r"^by no later than", "Once", None, None),
    (r"^\(a\) (start:|at least 5 business days)", "Once", None, None),
    (r"^\(?[ab]\) as soon as reasonably practicable", "Once", None, None),
    (r"^\(a\) as stated in the marine plant restoration plan", "Ongoing", "Restoration works completion", None),
    (r"^\(?b\) ?\(i\) as stated", "Milestone", "Restoration works completion", None),
]

# Sheets without a timing column: default Ongoing, Once for pre-construction rows, plus these (D-04).
FREQUENCY_OVERRIDE = {
    ("PFAS CMP Requirements", 11): "Six-monthly",  # "six-monthly review of the quantities of water"
    ("PFAS CMP Requirements", 12): "Daily",        # "Daily checks of the treatment system"
}

# Milestone obligations get Y1/Y2/Y5 occurrences, anchored by stage.
MILESTONE_YEARS = [("Y1", "Year 1", 365), ("Y2", "Year 2", 730), ("Y5", "Year 5", 1825)]

# Rows combining a commencement and a completion notice (D-05): split into two occurrences.
NOTIFICATION_SPLIT = {
    ("Bridge Creek DA Conditions", 10),
    ("Marine Plants DA Conditions", 24),
    ("Marine Plants DA Conditions", 39),
}

# CEMP row 11: two deliverables in one row (D-05).
CEMP_SPLIT_ROW = 11

ROLE_TOKENS = [  # (prefix of lower-case token, role or None for supporting-only, internal?)
    ("environmental and sustainability team", "Environmental and Sustainability Team", True),
    ("environmental manager", "Environmental Manager", True),
    ("environmental advisor", "Environmental Advisor", True),
    ("project manager", "Project Manager", True),
    ("construction manager", "Construction Manager", True),
    ("site manager", "Site Manager", True),
    ("site supervisor", "Site Supervisor", True),
    ("engineering and design team", "Engineering and Design Team", True),
    ("dt infrastructure", "Project Manager", True),              # D-11
    ("contractor", "Contractor", False),
    ("asset owner", "Asset Owner / Operator", False),
    ("maintenance manager", "Asset Owner / Operator", False),
    ("all dti workers", None, False),
    ("ecologist", None, False),
    ("fauna spotter-catcher", None, False),
    ("air quality consultant", None, False),
    ("environmental consultant", None, False),
    ("suitably qualified person", None, False),
]
DEFAULT_ROLE = "Environmental and Sustainability Team"  # D-11: blank responsibility

RISK_SORT = {"A": 1, "B": 2, "C": 3, "D": 4}
DAF_QUOTE_PREFIX = "the department believes"

# ---------------------------------------------------------------------------
# Source documents (D-14): _Doc Ref rows plus permits and plans the register cites
# ---------------------------------------------------------------------------
DOC_REF_KEYS = {  # _Doc Ref row -> DocKey
    5: "SCHED-C1", 6: "MD-13-320", 7: "MD-12-164", 8: "MD-15-3", 9: "MD-15-315", 10: "MD-15-316",
    11: "MD-15-317", 12: "ENV-PROCESSES-MANUAL", 13: "SCHED-A43", 14: "SCHED-A44", 15: "ADR-WWBW",
    16: "RPP-WSS-2013-726", 17: "CMP-OMF",
}
DOC_CATEGORY = {"Specification": "Specification", "Heritage": "Heritage", "WWBW": "WWBW",
                "RPP Exemption": "RPP Exemption", "CMP": "CMP"}
EXTRA_DOCS = [  # DocKey, title, category, reference
    ("SDA-2402-39300", "SARA Decision Notice 2402-39300 SDA - Bridge Creek Diversion", "Development approval", "2402-39300 SDA"),
    ("SDA-2506-46598", "SARA Decision Notice 2506-46598 SDA - No Name Creek waterway barrier works", "Development approval", "2506-46598 SDA"),
    ("SDA-2308-36075", "SARA Decision Notice 2308-36075 SDA - Marine plants (TNAR, temporary)", "Development approval", "2308-36075 SDA"),
    ("SDA-2309-36966", "SARA Decision Notice 2309-36966 SDA - Marine plants (NAC, permanent)", "Development approval", "2309-36966 SDA"),
    ("EMP", "Environmental Management Plan - Ormeau Facility", "Management plan", "EMP (D)"),
    ("CEMP", "Construction Environmental Management Plan - Ormeau Facility (incl. annexes)", "Management plan", "CEMP (B)"),
    ("PLAN-FISH-SALVAGE", "Fish Salvage Plan - Bridge Creek", "Management plan", "Fish Salvage Plan"),
    ("PLAN-REHAB", "Rehabilitation & Remediation Plan - Bridge Creek", "Management plan", "Rehabilitation & Remediation Plan"),
    ("PLAN-FISH-PASSAGE", "Fish Passage Monitoring Plan - Bridge Creek", "Management plan", "Fish Passage Monitoring Plan"),
    ("PLAN-AQUATIC-PLANT", "Aquatic Plant Monitoring Plan - Bridge Creek", "Management plan", "Aquatic Plant Monitoring Plan"),
    ("PLAN-MARINE-RESTORATION", "Marine Plant Restoration Plan (WGJV)", "Management plan", "Marine Plant Restoration Plan"),
    ("SUBCONTRACT-A", "Subcontract for Ormeau Facility - Part A Preliminary Matters", "Contract", "Subcontract Part A"),
    ("SCHED-A5", "Schedule A5 Management Systems and Management Plans", "Specification", "Schedule A5"),
    ("ISO-14001", "ISO 14001:2015 Environmental management systems", "Standard", "ISO 14001:2015"),
]
# Group-row and folder links for the permits (display text -> DocKey, field)
PERMIT_BY_BLOCK = {"TNAR": "SDA-2308-36075", "NAC": "SDA-2309-36966"}

# Source Reference text -> DocKey (first match wins)
SOURCE_DOC_RULES = [
    (r"^cmp\b|cmp ormeau facility", "CMP-OMF"),
    (r"^emp \(d\)", "EMP"),
    (r"^cemp", "CEMP"),
    (r"iso 14001", "ISO-14001"),
    (r"^accepted development requirements", "ADR-WWBW"),
    (r"^fish salvage plan", "PLAN-FISH-SALVAGE"),
    (r"^rehabilitation & remediation plan", "PLAN-REHAB"),
    (r"^fish passage monitoring plan", "PLAN-FISH-PASSAGE"),
    (r"^aquatic plant monitoring plan", "PLAN-AQUATIC-PLANT"),
    (r"cultural heritage management plan - jabree", "SCHED-A43"),
    (r"partial clearance report - danggan balun", "SCHED-A44"),
    (r"riverine protection permit exemption", "RPP-WSS-2013-726"),
    (r"^subcontract for ormeau facility", "SUBCONTRACT-A"),
    (r"^schedule a5", "SCHED-A5"),
    (r"^schedule c1", "SCHED-C1"),
    (r"^md-13-320", "MD-13-320"),
    (r"^md-15-3 ", "MD-15-3"),
]
SHEET_DEFAULT_DOC = {
    "Bridge Creek DA Conditions": "SDA-2402-39300",
    "No Name DA Conditions": "SDA-2506-46598",
}

# ---------------------------------------------------------------------------
# Live fields: set on create and filled when blank in SharePoint, but never
# overwritten on re-run (the team owns them after go-live) unless the loader
# is run with -OverwriteLiveFields.
# ---------------------------------------------------------------------------
LIST_META = {
    "SourceDocuments": {"keyFields": ["DocKey"], "liveFields": ["Version", "DocDate", "DocLink", "FolderLink", "Provided", "Comments"]},
    "ProjectMilestones": {"keyFields": ["Title"], "liveFields": ["PlannedDate", "ActualDate", "Notes"]},
    "RoleAssignments": {"keyFields": ["Title"], "liveFields": ["PrimaryPerson", "AdditionalPeople", "Notes"]},
    "Obligations": {"keyFields": ["Title"], "liveFields": [
        "ProjectStage", "RiskRating", "RiskSort", "Frequency", "RecurrenceStart", "RecurrenceEnd", "DueDate",
        "Owner", "ResponsibleRole", "Status", "StatusComment", "ComplianceOutcome", "ComplianceNotes",
        "LastReviewed", "Active", "HasOpenFollowUp", "IsOverdue", "CurrentOccurrenceKey"]},
    "Occurrences": {"keyFields": ["Title"], "liveFields": [
        "PeriodStart", "PeriodEnd", "PeriodDue", "Assignee", "Status", "StatusComment", "ComplianceOutcome",
        "CompletedDate", "EvidenceURL", "Notes", "IsOverdue", "HasOpenFollowUp"]},
    "Tasks": {"keyFields": ["ObligationID", "TaskOrder"], "liveFields": ["Assignee", "DueDate", "Done", "DoneDate", "Notes"]},
    "Evidence": {"keyFields": ["ObligationID", "OccurrenceKey", "EvidenceURL"], "liveFields": ["EvidenceDate", "AddedBy", "Notes"]},
    "FollowUps": {"keyFields": ["ObligationID", "Title"], "liveFields": [
        "Owner", "FollowUpDate", "Status", "ResolutionComment", "ResolvedBy", "ResolvedDate"]},
    "StatusHistory": {"keyFields": ["Title", "ItemType", "ChangeSource"], "createOnly": True, "liveFields": []},
    "ReportComponents": {"keyFields": ["Title", "ReportName"], "liveFields": []},
}
LOAD_ORDER = ["SourceDocuments", "ProjectMilestones", "RoleAssignments", "Obligations", "Occurrences",
              "Tasks", "Evidence", "FollowUps", "StatusHistory", "ReportComponents"]


# ---------------------------------------------------------------------------
# Helpers
# ---------------------------------------------------------------------------
def norm(text: Any) -> str:
    """Lower-case, whitespace-collapsed, citation-free text for rule matching."""
    if text is None:
        return ""
    t, _ = L.strip_citations(L.clean_text(str(text)) or "")
    return " ".join((t or "").split()).lower()


def one_line(text: str | None, limit: int = 250) -> str | None:
    if not text:
        return None
    flat = " ".join(text.split())
    return flat if len(flat) <= limit else flat[: limit - 1].rstrip() + "…"


def iso(d: dt.date | dt.datetime | None) -> str | None:
    if d is None:
        return None
    return (d.date() if isinstance(d, dt.datetime) else d).isoformat()


def lookup(list_key: str, key: str | None) -> dict | None:
    return {"lookup": list_key, "key": key} if key else None


def url_value(url: str | None, description: str | None = None) -> dict | None:
    return {"url": url, "description": (description or url)[:255]} if url else None


def parse_date(text: str) -> str | None:
    text = (text or "").strip()
    for fmt in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return dt.datetime.strptime(text, fmt).date().isoformat()
        except ValueError:
            pass
    return None


@dataclass
class Report:
    unowned: list[tuple[str, str]] = field(default_factory=list)
    evidence_gaps: list[tuple[str, str, str]] = field(default_factory=list)
    comments: list[tuple[str, str, str, str]] = field(default_factory=list)
    followups: list[tuple[str, str]] = field(default_factory=list)
    undated: list[tuple[str, str]] = field(default_factory=list)
    unmatched_timing: list[tuple[str, str]] = field(default_factory=list)
    unmatched_status: list[tuple[str, str]] = field(default_factory=list)
    complete_without_evidence: list[str] = field(default_factory=list)
    notes: list[str] = field(default_factory=list)


@dataclass
class ParsedRow:
    sheet: str
    row: int
    block: L.Block
    cls: str
    values: dict[str, Any]           # cleaned values by canonical field
    raw: dict[str, Any]              # raw cell values
    links: list[tuple[str, str, str]]  # (cell, display text, url) for hyperlinked cells
    comments: list[tuple[str, str, str]]  # (cell, author, text)
    extra: list[tuple[str, str]] = field(default_factory=list)
    obligation_id: str = ""
    parent_id: str = ""


# ---------------------------------------------------------------------------
# Transform
# ---------------------------------------------------------------------------
class Transformer:
    def __init__(self, xlsx: Path, import_date: dt.date):
        warnings.simplefilter("ignore")
        self.xlsx = xlsx
        self.wb = openpyxl.load_workbook(xlsx)
        self.threaded = L.load_threaded_comments(xlsx)
        self.import_date = import_date
        self.report = Report()
        self.roles = self._load_roles()
        self.milestones = self._load_milestones()
        self.lists: dict[str, list[dict]] = {k: [] for k in LOAD_ORDER}
        self.row_to_id: dict[tuple[str, int], str] = {}

    # -- config -------------------------------------------------------------
    def _load_roles(self) -> dict[str, dict]:
        roles = {}
        with open(CONFIG / "role-assignments.csv", newline="", encoding="utf-8") as fh:
            for rec in csv.DictReader(fh):
                role = rec["Responsible Role"].strip()
                if role not in SCHEMA["choiceSets"]["ResponsibleRole"]:
                    raise ValueError(f"role-assignments.csv: unknown role {role!r}")
                roles[role] = {
                    "primary": rec["Primary Person"].strip() or None,
                    "additional": [p.strip() for p in re.split(r"[;,]", rec["Additional People"]) if p.strip()],
                    "notes": rec["Notes"].strip() or None,
                }
        return roles

    def _load_milestones(self) -> dict[str, dict]:
        out = {}
        with open(CONFIG / "project-milestones.csv", newline="", encoding="utf-8") as fh:
            for rec in csv.DictReader(fh):
                name = rec["Milestone"].strip()
                if name not in SCHEMA["choiceSets"]["DueAnchor"]:
                    raise ValueError(f"project-milestones.csv: unknown milestone {name!r}")
                out[name] = {"planned": parse_date(rec["Planned Date"]), "actual": parse_date(rec["Actual Date"]),
                             "notes": rec["Notes"].strip() or None}
        return out

    def milestone_date(self, anchor: str | None, offset: int | None) -> str | None:
        if not anchor or offset is None:
            return None
        m = self.milestones.get(anchor) or {}
        base = m.get("actual") or m.get("planned")
        return (dt.date.fromisoformat(base) + dt.timedelta(days=offset)).isoformat() if base else None

    # -- parsing ------------------------------------------------------------
    def parse_rows(self) -> list[ParsedRow]:
        rows: list[ParsedRow] = []
        for spec in L.SHEETS:
            if spec.kind != "obligations":
                continue
            ws = self.wb[spec.name]
            for block in spec.blocks:
                header_texts = {str(c.value).strip() for c in ws[block.header_row] if c.value is not None}
                parent: ParsedRow | None = None
                block_location = None
                for r in range(block.first_row, block.last_row + 1):
                    cls, _ = L.classify(ws, r, block, header_texts)
                    if cls not in ("obligation", "info-only", "sub-condition", "flag"):
                        continue
                    raw = L.row_values(ws, r, block.columns)
                    values = {k: L.clean_text(v) for k, v in raw.items()}
                    links, comments = [], []
                    for col in block.columns:
                        cell = ws[f"{col}{r}"]
                        if cell.hyperlink:
                            target = cell.hyperlink.target or cell.hyperlink.location
                            links.append((cell.coordinate, L.clean_text(cell.value) or "Link", L.resolve_hyperlink(target)))
                    for (sheet, ref), thread in self.threaded.items():
                        if sheet == spec.name and int(re.sub(r"[A-Z]+", "", ref)) == r:
                            comments.extend((ref, a, t) for a, t in thread)
                    pr = ParsedRow(spec.name, r, block, cls, values, raw, links, comments)
                    self._fix_row(pr)
                    # Marine Plants: "Same project" and sub-conditions inherit from the block / parent row.
                    loc = pr.values.get("location")
                    if loc and loc.lower() == "same project":
                        pr.extra.append(("TP002 Project Location/Activity", loc))
                        pr.values["location"] = block_location
                    elif loc and block_location is None:
                        block_location = loc
                    if cls == "sub-condition" and parent is not None:
                        for k in ("location", "stage", "responsibility"):
                            if not pr.values.get(k):
                                pr.values[k] = parent.values.get(k)
                        pr.parent_id = "pending"
                        pr.extra.append(("Inherited from parent row", f"{spec.name}!row {parent.row}"))
                    else:
                        parent = pr
                    pr._parent = parent  # type: ignore[attr-defined]
                    rows.append(pr)
        return rows

    def _fix_row(self, pr: ParsedRow) -> None:
        v = pr.values
        # Marine Plants row 23: status shifted into Notes (docs/data-audit.md 2.3, D-08).
        if ("condition_no" in pr.block.columns.values() and norm(v.get("notes")) in STATUS_MAP and norm(v.get("notes"))
                and norm(v.get("status")) in STATUS_MAP and norm(v.get("status")) != norm(v.get("notes"))):
            pr.extra.append(("Unmapped (J)", v["status"]))
            v["status"], v["notes"], v["evidence"] = v["notes"], v.get("evidence"), None
            self.report.notes.append(f"{pr.sheet} row {pr.row}: status/notes realigned (row shifted one column right).")
        # Bridge Creek pasted citations (D-09).
        for key in ("requirement", "source_ref", "timing"):
            text, cites = L.strip_citations(v.get(key))
            if cites:
                v[key] = text
                pr.extra.append((f"Citations ({key})", " ".join(cites)))
        # DAF letter quote in the Evidence or unheaded Links column -> notes.
        for key in ("evidence", "link1", "link2", "link3"):
            text = v.get(key)
            if text and text.lower().startswith(DAF_QUOTE_PREFIX):
                v["notes"] = "\n".join(filter(None, [v.get("notes"), f"DAF correspondence: {text}"]))
                v[key] = None
        # Non-hyperlinked Links cells are placeholders / evidence gaps.
        linked_cells = {c for c, _, _ in pr.links}
        for col, fld in pr.block.columns.items():
            if fld.startswith("link") and v.get(fld) and f"{col}{pr.row}" not in linked_cells:
                pr.extra.append(("Links (not hyperlinked)", v[fld]))
                self.report.evidence_gaps.append(("", f"{pr.sheet}!{col}{pr.row}", v[fld]))
        # Editorial note in ADR location (row 6).
        if pr.sheet == "ADR WWBW Conditions" and (v.get("location") or "").lower().startswith("to change all of these"):
            pr.extra.append(("Editorial note (TP002 location)", v["location"]))
            v["location"] = None
            pr._editorial = True  # type: ignore[attr-defined]

    # -- IDs ------------------------------------------------------------------
    def assign_ids(self, rows: list[ParsedRow]) -> None:
        counters: Counter = Counter()
        for pr in rows:
            prefix = PREFIX.get(pr.sheet) or f"MP-{pr.block.context['permit_short']}"
            if pr.cls == "sub-condition":
                parent = pr._parent  # type: ignore[attr-defined]
                suffix = (re.sub(r"[^a-z]", "", (pr.values.get("condition_no") or "").lower())[-1:] or "x").upper()
                pr.obligation_id = f"{parent.obligation_id}{suffix}"
                pr.parent_id = parent.obligation_id
            else:
                counters[prefix] += 1
                pr.obligation_id = f"{prefix}-{counters[prefix]:03d}"
            self.row_to_id[(pr.sheet, pr.row)] = pr.obligation_id

    # -- roles --------------------------------------------------------------
    def resolve_roles(self, responsibility: str | None) -> tuple[str | None, str | None]:
        if not responsibility:
            return DEFAULT_ROLE, None
        tokens = [t.strip() for t in re.split(r"\n| / ", responsibility) if t.strip()]
        internal, external, supporting = [], [], []
        for tok in tokens:
            low = tok.lower()
            match = next((m for m in ROLE_TOKENS if low.startswith(m[0])), None)
            if match is None:
                supporting.append(tok)
                self.report.notes.append(f"Unrecognised responsibility token {tok!r} kept in Supporting Roles.")
            elif match[1] is None:
                supporting.append(tok)
            elif match[2]:
                internal.append((match[1], tok))
            else:
                external.append((match[1], tok))
        ranked = internal + external
        primary = ranked[0][0] if ranked else DEFAULT_ROLE
        rest = [orig for role, orig in ranked[1:]] + supporting
        seen, dedup = set(), []
        for r in rest:
            if r.lower() not in seen:
                seen.add(r.lower())
                dedup.append(r)
        return primary, ("; ".join(dedup) or None)

    # -- frequency ----------------------------------------------------------
    def frequency(self, pr: ParsedRow, stage: str) -> tuple[str, str | None, int | None]:
        key = (pr.sheet, pr.row)
        if key in FREQUENCY_OVERRIDE:
            return FREQUENCY_OVERRIDE[key], None, None
        timing = norm(pr.values.get("timing"))
        if timing:
            for pattern, freq, anchor, offset in TIMING_RULES:
                if re.search(pattern, timing):
                    return freq, anchor, offset
            self.report.unmatched_timing.append((pr.obligation_id, pr.values.get("timing") or ""))
            return "Ongoing", None, None
        if stage == "Pre-construction":
            return "Once", "Construction start", 0
        return "Ongoing", None, None

    # -- main build -----------------------------------------------------------
    def build(self) -> None:
        # No Name row 7 lists works (a)-(d) under condition 1: imported as Tasks, not an obligation (D-06).
        rows = [r for r in self.parse_rows() if (r.sheet, r.row) != ("No Name DA Conditions", 7)]
        self.assign_ids(rows)
        self.build_source_documents()
        self.build_milestones()
        self.build_roles()
        for pr in rows:
            self.build_obligation(pr)
        self.build_report_components()

    def _extra_fields(self, pr: ParsedRow, originals: dict[str, str | None]) -> str:
        cols = list(pr.block.columns)
        parts = [("Source cell range", f"'{pr.sheet}'!{cols[0]}{pr.row}:{cols[-1]}{pr.row}")]
        if pr.block.context.get("permit"):
            parts.append(("Permit / decision notice", pr.block.context["permit"]))
        for label, value in originals.items():
            if value:
                parts.append((label, value))
        if pr.values.get("legislation"):
            parts.append(("Relevant Legislation / Standard or Code", pr.values["legislation"]))
        parts.extend(pr.extra)
        for cell, author, text in pr.comments:
            parts.append((f"Excel comment ({cell}, {author})", text))
            self.report.comments.append((pr.obligation_id, cell, author, text))
        return "\n".join(f"{label}: {value}" for label, value in parts)

    def build_obligation(self, pr: ParsedRow) -> None:
        v = pr.values
        oid = pr.obligation_id
        info_only = pr.cls == "info-only"
        raw_stage = v.get("stage")
        stage = STAGE_MAP.get(norm(raw_stage), "All stages") if raw_stage else "All stages"
        if raw_stage and norm(raw_stage) not in STAGE_MAP:
            self.report.notes.append(f"{oid}: unknown stage {raw_stage!r} -> All stages")

        # Condition number: explicit column, or "Condition 12" in Bridge Creek's Source/Reference.
        condition = v.get("condition_no")
        condition = str(condition) if condition is not None else None
        if not condition and pr.sheet == "Bridge Creek DA Conditions":
            m = re.match(r"condition (\d+)", norm(v.get("source_ref")))
            condition = m.group(1) if m else None

        # Section/group: first line kept, rest recorded (CEMP "General\nLegal Compliance").
        section = v.get("section")
        section_extra = None
        if section and "\n" in section:
            section_extra = section.replace("\n", " / ")
            section = section.split("\n")[0]
        if not section and pr.block.context.get("permit_short"):
            section = f"{pr.block.context['permit_short']} ({pr.block.context['permit']})"

        # Requirement
        requirement = v.get("requirement")
        if pr.cls == "flag":
            if requirement and len(requirement) < 5:
                pr.extra.append(("TP002 Requirement (as entered)", requirement))
                requirement = f"⚠ Requirement text missing in TP002 {EN_DASH} see source reference"
            elif not requirement:
                requirement = v.get("source_ref") or "⚠ Requirement text missing in TP002"

        # Roles and owner
        if info_only:
            role, supporting = None, None
        else:
            role, supporting = self.resolve_roles(v.get("responsibility"))
        owner = self.roles.get(role, {}).get("primary") if role else None
        if role and not owner:
            self.report.unowned.append((oid, role))

        # Frequency / due
        freq, anchor, offset = ("Ongoing", None, None) if info_only else self.frequency(pr, stage)
        split_cemp = pr.sheet == "CEMP and General Requirements" and pr.row == CEMP_SPLIT_ROW
        due_raw = v.get("due")
        due_date = iso(due_raw) if isinstance(due_raw, (dt.date, dt.datetime)) else None
        due_rule = None
        if isinstance(due_raw, (dt.date, dt.datetime)):
            due_rule = one_line(v.get("timing"), 250)
        elif due_raw:
            due_rule = one_line(str(due_raw), 250)
        if not due_date and anchor and offset is not None and freq == "Once":
            due_date = self.milestone_date(anchor, offset)
            if not due_date:
                self.report.undated.append((oid, f"{anchor} + {offset} days"))
        due_offset = offset if freq == "Once" else None
        if pr.sheet == "CEMP and General Requirements" and "5th day" in norm(v.get("timing")):
            due_offset = 5  # monthly report due by the 5th of the following month

        # Status
        raw_status = v.get("status") or ""
        if info_only:
            status = NA_INFO
        elif (pr.sheet, pr.row) in STATUS_OVERRIDE:
            status = STATUS_OVERRIDE[(pr.sheet, pr.row)]
        else:
            status = STATUS_MAP.get(norm(raw_status))
            if status is None:
                self.report.unmatched_status.append((oid, raw_status))
                status = "Not started"
        if status == NA_INFO and not info_only:
            info_only = True

        risk = v.get("risk") if v.get("risk") in RISK_SORT else None
        src_doc = self.source_doc_for(pr, v.get("source_ref"))

        originals = {
            "TP002 Project Stage": raw_stage if raw_stage and raw_stage != stage else None,
            "TP002 Compliance Status": raw_status or None,
            "TP002 Responsibility": (v.get("responsibility") or "").replace("\n", " / ") or None,
            "Section/Group (full)": section_extra,
        }
        evidence_required = v.get("evidence") or v.get("compliance_check")

        base = {
            "Title": oid,
            "SourceRegister": pr.sheet,
            "SourceOrder": pr.row,
            "SourceCell": f"'{pr.sheet}'!A{pr.row}",
            "SectionGroup": one_line(section, 255),
            "ConditionNo": condition,
            "ParentObligationID": pr.parent_id or None,
            "LocationActivity": one_line(v.get("location"), 255),
            "ProjectStage": stage,
            "RequirementSummary": one_line(requirement),
            "Requirement": requirement,
            "SourceReference": v.get("source_ref"),
            "SourceDocument": lookup("SourceDocuments", src_doc),
            "ComplianceAction": v.get("action"),
            "PracticalImplementation": v.get("practical_impl"),
            "EvidenceRequired": evidence_required,
            "Stakeholder": v.get("stakeholder"),
            "RiskRating": risk,
            "RiskSort": RISK_SORT.get(risk, 5),
            "Frequency": freq,
            "TimingDetail": v.get("timing"),
            "RecurrenceStart": self.milestone_date(anchor, 0) if freq in ("Quarterly",) and anchor else None,
            "RecurrenceEnd": None,
            "DueAnchor": anchor,
            "DueOffsetDays": due_offset,
            "DueDate": due_date,
            "DueRule": due_rule,
            "Owner": owner,
            "ResponsibleRole": role,
            "SupportingRoles": supporting,
            "Status": status,
            "StatusComment": None,
            "ComplianceOutcome": "Not assessed",
            "ComplianceNotes": v.get("notes"),
            "LastReviewed": iso(v["last_reviewed"]) if isinstance(v.get("last_reviewed"), (dt.date, dt.datetime)) else None,
            "InfoOnly": info_only,
            "Active": not info_only,
            "HasOpenFollowUp": False,
            "IsOverdue": False,
            "CurrentOccurrenceKey": None,
            "ExtraFields": None,
        }

        items = [base]
        if split_cemp:
            # Weekly inspection checklist (base) + monthly CMP audit (B) - D-05.
            actions = [a for a in (v.get("action") or "").split("\n") if a.strip()]
            second = dict(base)
            base["ComplianceAction"] = actions[0] if actions else base["ComplianceAction"]
            base["Frequency"] = "Weekly"
            second.update({"Title": f"{oid}B", "ComplianceAction": actions[1] if len(actions) > 1 else None,
                           "Frequency": "Monthly", "ParentObligationID": oid})
            pr.extra.append(("Split", f"Row split into {oid} (weekly inspection) and {oid}B (monthly CMP audit) - D-05"))
            items.append(second)
            self.row_to_id[(pr.sheet, pr.row, "monthly")] = f"{oid}B"  # type: ignore[index]

        extra = self._extra_fields(pr, originals)
        for item in items:
            item["ExtraFields"] = extra

        # Occurrences, evidence, tasks, follow-ups
        occurrences = self.build_occurrences(pr, base)
        self.build_evidence(pr, base, occurrences)
        self.build_followups(pr, items)

        for item in items:
            if item is base and occurrences:
                self.roll_up(item, occurrences)
            self.lists["Obligations"].append(item)
            self.history(item["Title"], "Obligation", item["Status"], item["Title"],
                         f"Initial status from TP002.1 Rev 2 (TP002 value: {raw_status or 'blank'})")
            if item["Status"] == "Complete":
                has_ev = any(e["ObligationID"] == item["Title"] for e in self.lists["Evidence"])
                if not has_ev:
                    self.report.complete_without_evidence.append(item["Title"])

    def source_doc_for(self, pr: ParsedRow, source_ref: str | None) -> str | None:
        if pr.block.context.get("permit_short"):
            return PERMIT_BY_BLOCK[pr.block.context["permit_short"]]
        text = norm(source_ref)
        for pattern, key in SOURCE_DOC_RULES:
            if re.search(pattern, text):
                return key
        return SHEET_DEFAULT_DOC.get(pr.sheet)

    # -- occurrences ----------------------------------------------------------
    def _occurrence(self, ob: dict, code: str, label: str, status: str, due: str | None,
                    evidence_url: dict | None = None, notes: str | None = None) -> dict:
        return {
            "Title": f"{ob['Title']}|{code}",
            "Obligation": lookup("Obligations", ob["Title"]),
            "ObligationID": ob["Title"],
            "ObligationSummary": ob["RequirementSummary"],
            "PeriodLabel": label,
            "PeriodStart": None,
            "PeriodEnd": None,
            "PeriodDue": due,
            "Assignee": ob["Owner"],
            "Status": status,
            "StatusComment": None,
            "ComplianceOutcome": "Not assessed",
            "CompletedDate": None,
            "EvidenceURL": evidence_url,
            "Notes": notes,
            "RiskSort": ob["RiskSort"],
            "IsOverdue": False,
            "HasOpenFollowUp": False,
        }

    def build_occurrences(self, pr: ParsedRow, ob: dict) -> list[dict]:
        occ: list[dict] = []
        if (pr.sheet, pr.row) in NOTIFICATION_SPLIT:
            commence_link = pr.links[0] if pr.links else None
            occ.append(self._occurrence(
                ob, "COMMENCE", "Commencement notice",
                "Complete" if commence_link else "Not started",
                self.milestone_date("Construction start", -5),
                url_value(commence_link[2], commence_link[1]) if commence_link else None,
                "Commencement notification evidenced in TP002 (link)." if commence_link else None))
            occ.append(self._occurrence(
                ob, "COMPLETE", "Completion notice", "Not started",
                self.milestone_date("Practical completion", 21),
                notes="Completion notification outstanding (TP002: 'Required at end of conditions')."))
            ob["Frequency"] = "Once"
        elif ob["Frequency"] == "Milestone" and ob["DueAnchor"] != "Restoration works completion":
            anchor = ob["DueAnchor"] or ("Rehabilitation completion" if norm(pr.values.get("stage")).startswith("rehabilitation")
                                         else "Practical completion")
            ob["DueAnchor"] = anchor
            for code, label, days in MILESTONE_YEARS:
                due = self.milestone_date(anchor, days)
                if not due:
                    self.report.undated.append((f"{ob['Title']}|{code}", f"{anchor} + {days} days"))
                occ.append(self._occurrence(ob, code, label, "Not started", due))
        elif ob["Frequency"] == "Milestone":
            self.report.undated.append((ob["Title"], "Milestones per the Marine Plant Restoration Plan - add occurrences once the plan schedule is known"))
        for o in occ:
            if o["Status"] == "Complete":
                o["CompletedDate"] = None
            self.history(o["Title"], "Occurrence", o["Status"], ob["Title"], "Created at import")
        self.lists["Occurrences"].extend(occ)
        return occ

    @staticmethod
    def roll_up(ob: dict, occ: list[dict]) -> None:
        """docs/list-schema.md section 3: status comes from the current occurrence."""
        statuses = [o["Status"] for o in occ]
        open_occ = [o for o in occ if o["Status"] != "Complete"]
        current = open_occ[0] if open_occ else occ[-1]
        ob["CurrentOccurrenceKey"] = current["Title"]
        if current["PeriodDue"]:
            ob["DueDate"] = current["PeriodDue"]
        if all(s == "Complete" for s in statuses):
            ob["Status"] = "Complete"
        elif "Missing" in statuses:
            ob["Status"] = "Missing"
        elif current["Status"] == "Not started" and "Complete" in statuses:
            ob["Status"] = "Underway"
        else:
            ob["Status"] = current["Status"]

    # -- evidence -------------------------------------------------------------
    def build_evidence(self, pr: ParsedRow, ob: dict, occ: list[dict]) -> None:
        commence_key = next((o["Title"] for o in occ if o["Title"].endswith("|COMMENCE")), None)
        seen = set()
        for i, (cell, text, url) in enumerate(pr.links):
            occ_key = commence_key if (commence_key and i == 0) else None
            if (url, occ_key) in seen:
                continue
            seen.add((url, occ_key))
            self.lists["Evidence"].append({
                "Title": one_line(text, 255),
                "ObligationID": ob["Title"],
                "OccurrenceKey": occ_key,
                "EvidenceURL": url_value(url, text),
                "EvidenceDate": None,
                "AddedBy": None,
                "Notes": f"Imported from TP002 {pr.sheet}!{cell}",
            })

    # -- follow-ups (D-12) ------------------------------------------------------
    def build_followups(self, pr: ParsedRow, items: list[dict]) -> None:
        reasons = []
        if pr.cls == "flag":
            reasons.append(("Requirement text missing or incomplete in TP002",
                            "The TP002 row has no usable requirement text. Confirm the obligation from the source document "
                            "and update the Requirement."))
        if getattr(pr, "_editorial", False):
            reasons.append(("TP002 editorial note to resolve",
                            f"TP002 location column contained an editorial note: '{[e for e in pr.extra if e[0].startswith('Editorial')][0][1]}'. "
                            "Confirm the creek crossings that required ADR WWBW and update the location."))
        if pr.sheet == "CEMP and General Requirements" and pr.row == 41:
            reasons.append(("Confirm review frequency",
                            "TP002 timing was 'Quarterly / Six monthly?'. Imported as Quarterly; confirm and update Frequency."))
        if pr.sheet == "CEMP and General Requirements" and pr.row == 28:
            reasons.append(("Stage and timing conflict",
                            "TP002 stage is Pre construction but timing is Post construction. Confirm when this applies."))
        ob = items[0]
        for title, reason in reasons:
            self.lists["FollowUps"].append({
                "Title": title,
                "ObligationID": ob["Title"],
                "OccurrenceKey": None,
                "ObligationSummary": ob["RequirementSummary"],
                "ReasonType": "Import review",
                "Reason": reason,
                "Owner": ob["Owner"],
                "FollowUpDate": (self.import_date + dt.timedelta(days=14)).isoformat(),
                "Status": "Open",
                "ResolutionComment": None,
                "ResolvedBy": None,
                "ResolvedDate": None,
                "RiskSort": ob["RiskSort"],
            })
            ob["HasOpenFollowUp"] = True
            self.report.followups.append((ob["Title"], title))

    # -- status history -------------------------------------------------------
    def history(self, item_key: str, item_type: str, status: str, obligation_id: str, comment: str) -> None:
        self.lists["StatusHistory"].append({
            "Title": item_key,
            "ItemType": item_type,
            "ObligationID": obligation_id,
            "OldStatus": None,
            "NewStatus": status,
            "ChangedBy": None,
            "ChangedAt": f"{self.import_date.isoformat()}T09:00:00+10:00",
            "Comment": comment,
            "ChangeSource": "Import",
        })

    # -- reference lists ------------------------------------------------------
    def build_source_documents(self) -> None:
        ws = self.wb["_Doc Ref"]
        for r, key in DOC_REF_KEYS.items():
            title = L.clean_text(ws[f"B{r}"].value)
            font = ws[f"B{r}"].font
            provided = not (font and font.color is not None and font.color.rgb == "FFFF0000")
            version = ws[f"D{r}"].value
            date = ws[f"E{r}"].value
            self.lists["SourceDocuments"].append({
                "DocKey": key,
                "Title": one_line(title, 255),
                "Category": DOC_CATEGORY.get(L.clean_text(ws[f"A{r}"].value)),
                "Reference": L.clean_text(str(ws[f"C{r}"].value)) if ws[f"C{r}"].value is not None else None,
                "Version": str(version) if version is not None else None,
                "DocDate": iso(date) if isinstance(date, (dt.date, dt.datetime)) else None,
                "DocLink": None,
                "FolderLink": None,
                "Provided": provided,
                "Comments": L.clean_text(ws[f"F{r}"].value),
            })
        mp = self.wb["Marine Plants DA Conditions"]
        permit_links = {
            "SDA-2308-36075": ("B21", "C21", "J6"),
            "SDA-2309-36966": ("B32", "C32", "J7"),
        }
        folder_links = {"SDA-2402-39300": ("Bridge Creek DA Conditions", "L6"),
                        "SDA-2506-46598": ("No Name DA Conditions", "M6")}
        for key, title, category, ref in EXTRA_DOCS:
            doc_link = folder_link = None
            comments = "Added at import: cited in the register but not listed on TP002 _Doc Ref (D-14)."
            if key in permit_links:
                b, c, legacy = permit_links[key]
                doc_link = url_value(L.resolve_hyperlink(mp[b].hyperlink.target), one_line(L.clean_text(mp[b].value), 255))
                folder_link = url_value(L.resolve_hyperlink(mp[c].hyperlink.target), "Decision documents")
                legacy_url = L.resolve_hyperlink(mp[legacy].hyperlink.target)
                legacy_row = int(legacy[1:])
                comments += (f"\nLegacy TP002 approval summary (hidden row {legacy_row}): approval link {legacy_url}; "
                             f"due {iso(mp[f'G{legacy_row}'].value)}; status {mp[f'H{legacy_row}'].value}.")
                if key == "SDA-2309-36966":
                    comments += " TP002 row 7 quoted this permit as 2309-36996 SDA; 2309-36966 matches the decision notice (D-08)."
            if key in folder_links:
                sheet, cell = folder_links[key]
                c = self.wb[sheet][cell]
                folder_link = url_value(L.resolve_hyperlink(c.hyperlink.target), one_line(L.clean_text(c.value), 255))
            self.lists["SourceDocuments"].append({
                "DocKey": key, "Title": title, "Category": category, "Reference": ref, "Version": None,
                "DocDate": None, "DocLink": doc_link, "FolderLink": folder_link, "Provided": True, "Comments": comments,
            })

    def build_milestones(self) -> None:
        for name, m in self.milestones.items():
            self.lists["ProjectMilestones"].append({
                "Title": name, "PlannedDate": m["planned"], "ActualDate": m["actual"], "Notes": m["notes"]})

    def build_roles(self) -> None:
        for role, r in self.roles.items():
            self.lists["RoleAssignments"].append({
                "Title": role, "PrimaryPerson": r["primary"], "AdditionalPeople": r["additional"] or None, "Notes": r["notes"]})

    def build_report_components(self) -> None:
        ws = self.wb["Reporting Matrix"]
        cemp = "CEMP and General Requirements"
        producers = {
            "B": self.row_to_id[(cemp, 10)],
            "C": self.row_to_id[(cemp, 9)],
            "D": self.row_to_id[(cemp, CEMP_SPLIT_ROW, "monthly")],  # type: ignore[index]
            "E": self.row_to_id[(cemp, 44)],  # D-13
            "F": self.row_to_id[(cemp, 8)],
        }
        names = SCHEMA["choiceSets"]["ReportName"]
        for i, col in enumerate("BCDEF"):
            header = L.clean_text(ws[f"{col}4"].value) or ""
            freq_text = header.split("\n")[0].strip().lower()
            freq = "Monthly" if freq_text == "monthly" else "Post construction"
            src = " ".join(t for _, t in self.threaded.get(("Reporting Matrix", f"{col}4"), []))
            for r in range(5, ws.max_row + 1):
                if str(ws[f"{col}{r}"].value or "").strip().upper() != "X":
                    continue
                self.lists["ReportComponents"].append({
                    "Title": one_line(L.clean_text(ws[f"A{r}"].value), 255),
                    "ReportName": names[i],
                    "ReportFrequency": freq,
                    "ReportSourceRef": src or None,
                    "ProducingObligationID": producers[col],
                    "SortOrder": r,
                })

    # -- NN condition 1 (a)-(d) tasks (D-06) ---------------------------------
    def build_nn_tasks(self) -> None:
        ws = self.wb["No Name DA Conditions"]
        text = L.clean_text(ws["D7"].value) or ""
        status_text = norm(ws["J7"].value)
        done_letters = set(re.findall(r"\(([a-d])\)\s*-\s*completed", status_text))
        parent_id = self.row_to_id[("No Name DA Conditions", 6)]
        parent = next(o for o in self.lists["Obligations"] if o["Title"] == parent_id)
        items = re.split(r"\n(?=\([a-z]\))", text)
        for order, item in enumerate(items, 1):
            letter = item[1]
            self.lists["Tasks"].append({
                "Title": one_line(item, 255),
                "Obligation": lookup("Obligations", parent_id),
                "ObligationID": parent_id,
                "ObligationSummary": parent["RequirementSummary"],
                "TaskOrder": order,
                "Assignee": parent["Owner"],
                "DueDate": None,
                "Done": letter in done_letters,
                "DoneDate": None,
                "Notes": item if len(item) > 250 else None,
            })
        parent["ExtraFields"] += "\nChecklist: works (a)-(d) from TP002 row 7 imported as Obligation Tasks (D-06)."


def run(xlsx: Path, out: Path, import_date: dt.date) -> Transformer:
    t = Transformer(xlsx, import_date)
    t.build()
    t.build_nn_tasks()
    validate(t)
    write_outputs(t, out)
    return t


# ---------------------------------------------------------------------------
# Validation against the list schema
# ---------------------------------------------------------------------------
def validate(t: Transformer) -> None:
    schema_lists = {l["key"]: l for l in SCHEMA["lists"]}
    errors = []
    for key, items in t.lists.items():
        lst = schema_lists[key]
        fields = {f["name"]: f for f in lst["fields"]}
        fields["Title"] = {"name": "Title", "type": "Text", "required": True}
        for item in items:
            label = f"{key}:{item.get('Title') or item.get('DocKey')}"
            for name, value in item.items():
                if name not in fields:
                    errors.append(f"{label}: unknown column {name}")
                    continue
                f = fields[name]
                if f.get("required") and value in (None, ""):
                    errors.append(f"{label}: required column {name} is blank")
                if value is None:
                    continue
                if f["type"] == "Choice":
                    c = f["choices"]
                    opts = SCHEMA["choiceSets"][c[1:]] if isinstance(c, str) else c
                    if value not in opts:
                        errors.append(f"{label}: {name}={value!r} is not a valid choice")
                expected = {"Text": str, "Note": str, "Choice": str, "Number": (int, float), "Boolean": bool,
                            "DateOnly": str, "DateTime": str, "User": str, "UserMulti": list, "URL": dict, "Lookup": dict}[f["type"]]
                if not isinstance(value, expected) or (f["type"] == "Number" and isinstance(value, bool)):
                    errors.append(f"{label}: {name}={value!r} is not a {f['type']} value")
                elif f["type"] == "Text" and (len(value) > 255 or "\n" in value):
                    errors.append(f"{label}: {name} too long or multi-line for a single-line column")
                elif f["type"] == "DateOnly" and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", value):
                    errors.append(f"{label}: {name}={value!r} is not YYYY-MM-DD")
            for name, f in fields.items():
                if f.get("required") and name not in item:
                    errors.append(f"{label}: required column {name} missing")
        # Unique keys
        meta = LIST_META[key]
        keys = [tuple(str(i.get(k)) for k in meta["keyFields"]) for i in items]
        dupes = [k for k, n in Counter(keys).items() if n > 1]
        if dupes:
            errors.append(f"{key}: duplicate keys {dupes[:5]}")
    # Lookups resolve
    ob_ids = {o["Title"] for o in t.lists["Obligations"]}
    doc_ids = {d["DocKey"] for d in t.lists["SourceDocuments"]}
    for key, items in t.lists.items():
        for item in items:
            for name, value in item.items():
                if isinstance(value, dict) and "lookup" in value:
                    pool = ob_ids if value["lookup"] == "Obligations" else doc_ids
                    if value["key"] not in pool:
                        errors.append(f"{key}:{item.get('Title')}: lookup {name} -> {value['key']} not found")
    if errors:
        raise ValueError("Transform validation failed:\n  " + "\n  ".join(errors[:50]))


# ---------------------------------------------------------------------------
# Output
# ---------------------------------------------------------------------------
def flat(value: Any) -> Any:
    if isinstance(value, dict) and "lookup" in value:
        return value["key"]
    if isinstance(value, dict) and "url" in value:
        return value["url"]
    if isinstance(value, list):
        return "; ".join(map(str, value))
    return value


def write_outputs(t: Transformer, out: Path) -> None:
    out.mkdir(parents=True, exist_ok=True)
    schema_lists = {l["key"]: l for l in SCHEMA["lists"]}
    for key in LOAD_ORDER:
        items = t.lists[key]
        meta = LIST_META[key]
        doc = {"list": schema_lists[key]["title"], "url": schema_lists[key]["url"], "key": key,
               "keyFields": meta["keyFields"], "liveFields": meta["liveFields"],
               "createOnly": meta.get("createOnly", False), "count": len(items), "items": items}
        (out / f"{key}.json").write_text(json.dumps(doc, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
        cols = list(items[0].keys()) if items else []
        with open(out / f"{key}.csv", "w", newline="", encoding="utf-8-sig") as fh:
            w = csv.DictWriter(fh, fieldnames=cols)
            w.writeheader()
            for item in items:
                w.writerow({k: flat(v) for k, v in item.items()})
    (out / "import-report.md").write_text(render_report(t), encoding="utf-8")


def render_report(t: Transformer) -> str:
    r = t.report
    obs = t.lists["Obligations"]
    lines = ["# Import report", "",
             f"Generated by `import/transform.py` from `{t.xlsx.name}` (import date {t.import_date.strftime('%d/%m/%Y')}). "
             "Re-run the transform to refresh it.", "", "## Counts", "", "| List | Items |", "|---|---|"]
    for key in LOAD_ORDER:
        lines.append(f"| {key} | {len(t.lists[key])} |")
    by_reg = Counter(o["SourceRegister"] for o in obs)
    by_status = Counter(o["Status"] for o in obs)
    by_freq = Counter(o["Frequency"] for o in obs)
    lines += ["", "Obligations by register: " + ", ".join(f"{k} {v}" for k, v in by_reg.items()),
              "", "Obligations by status: " + ", ".join(f"{k} {v}" for k, v in by_status.most_common()),
              "", "Obligations by frequency: " + ", ".join(f"{k} {v}" for k, v in by_freq.most_common()), ""]

    def section(title: str, intro: str, header: str, rows: list[str]) -> None:
        lines.extend([f"## {title} ({len(rows)})", "", intro, ""])
        if rows:
            lines.extend([header, "|" + "---|" * (header.count("|") - 1)] + rows)
        else:
            lines.append("_None._")
        lines.append("")

    unowned_by_role = defaultdict(list)
    for oid, role in r.unowned:
        unowned_by_role[role].append(oid)
    section("Roles without a person", f"{len(r.unowned)} obligations have no owner yet. Their role has no Primary Person in `import/config/role-assignments.csv`. "
            "Add the person there (or in the Role Assignments list) and re-run the loader; blank owners are filled on re-run.",
            "| Role | Count | Obligations |",
            [f"| {role} | {len(ids)} | {', '.join(ids)} |" for role, ids in sorted(unowned_by_role.items())])
    section("Follow-ups created at import (D-12)", "Open follow-ups, owner = obligation owner, due 14 days after import.",
            "| Obligation | Follow-up |", [f"| {o} | {t_} |" for o, t_ in r.followups])
    section("Dates waiting on project milestones (D-10)",
            "Fill `import/config/project-milestones.csv` (or the Project Milestones list) and re-run to date these.",
            "| Item | Rule |", [f"| {o} | {rule} |" for o, rule in r.undated])
    gap_rows = []
    cell_to_id = {}
    for o in obs:
        m = re.search(r"'(.+)'!A(\d+)", o["SourceCell"])
        cell_to_id[(m.group(1), int(m.group(2)))] = o["Title"]
    for _, cell, text in r.evidence_gaps:
        sheet, ref = cell.rsplit("!", 1)
        oid = cell_to_id.get((sheet, int(re.sub(r"[A-Z]+", "", ref))), "")
        gap_rows.append(f"| {oid} | {cell} | {text} |")
    section("Evidence gaps (Links cells with no hyperlink)",
            "Kept in Extra Fields. '?', '..', 'drawings?' and 'monthly reports?' mean the evidence location is unknown; "
            "'N/A - until completion / no current records' mean none is expected yet.",
            "| Obligation | Cell | Text |", gap_rows)
    section("Excel review comments carried into Extra Fields", "Triage these and raise Follow-ups where action is needed.",
            "| Obligation | Cell | Author | Comment |",
            [f"| {o} | {c} | {a} | {x.replace('|', '/')} |" for o, c, a, x in r.comments])
    section("Complete without evidence", "Should be empty: Complete requires an evidence link or justification.",
            "| Obligation |", [f"| {o} |" for o in r.complete_without_evidence])
    section("Unmatched timing values (defaulted to Ongoing)", "Should be empty.", "| Obligation | Timing |",
            [f"| {o} | {x} |" for o, x in r.unmatched_timing])
    section("Unmatched status values (defaulted to Not started)", "Should be empty.", "| Obligation | Status |",
            [f"| {o} | {x} |" for o, x in r.unmatched_status])
    lines += ["## Notes", ""] + [f"- {n}" for n in r.notes] + [""]
    return "\n".join(lines)


def main() -> None:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    ap.add_argument("--xlsx", type=Path, default=DEFAULT_XLSX)
    ap.add_argument("--out", type=Path, default=DEFAULT_OUT)
    ap.add_argument("--import-date", type=dt.date.fromisoformat, default=dt.date.today())
    args = ap.parse_args()
    t = run(args.xlsx, args.out, args.import_date)
    print(f"Wrote {args.out}")
    for key in LOAD_ORDER:
        print(f"  {key:18} {len(t.lists[key]):4}")


if __name__ == "__main__":
    main()
